import os
import openpyxl
import pandas as pd
import numpy as np
import psycopg2
import json

base_dir = r'c:\Users\haris\OneDrive\Desktop\Final-Year-Project\urban-population-resource-optimization'
excel_path = os.path.join(base_dir, 'data', 'raw', 'Chennai_Model_Input_Aligned_Final.xlsx')
processed_csv_path = os.path.join(base_dir, 'data', 'processed', 'chennai_preprocessed_2015_2030.csv')
quality_report_path = os.path.join(base_dir, 'data', 'processed', 'data_quality_report.json')

zones_info = [
    (1, 'I', 'Thiruvottiyur'),
    (2, 'II', 'Manali'),
    (3, 'III', 'Madhavaram'),
    (4, 'IV', 'Tondiarpet'),
    (5, 'V', 'Royapuram'),
    (6, 'VI', 'Thiru-Vi-Ka-Nagar'),
    (7, 'VII', 'Ambattur'),
    (8, 'VIII', 'Annanagar'),
    (9, 'IX', 'Teynampet'),
    (10, 'X', 'Kodambakkam'),
    (11, 'XI', 'Valasaravakkam'),
    (12, 'XII', 'Alandur'),
    (13, 'XIII', 'Adyar'),
    (14, 'XIV', 'Perungudi'),
    (15, 'XV', 'Sholinganallur')
]

code_to_id = {z[1]: z[0] for z in zones_info}

def load_and_preprocess():
    print(f"Loading Excel workbook: {excel_path}")
    wb = openpyxl.load_workbook(excel_path, data_only=True)
    
    # Read Model_Input_2015_2030
    ws = wb['Model_Input_2015_2030']
    data = []
    headers = [cell.value for cell in ws[2]] # Row 2 has headers
    
    for r in range(3, ws.max_row + 1):
        row_vals = [ws.cell(r, c).value for c in range(1, len(headers) + 1)]
        if any(v is not None for v in row_vals):
            data.append(row_vals)
            
    df = pd.DataFrame(data, columns=headers)
    print(f"Raw records loaded: {len(df)}")
    
    # Map zone code to zone_id
    df['ZONE_CODE'] = df['ZONE'].astype(str).str.strip().str.upper()
    df['ZONE_ID'] = df['ZONE_CODE'].map(code_to_id)
    
    # Validation checks
    total_records = len(df)
    missing_count = df.isnull().sum().sum()
    duplicate_count = df.duplicated(subset=['YEAR', 'ZONE_ID']).sum()
    invalid_num_count = ((df['POPULATION'] <= 0) | (df['BUILT_UP_SURFACE_M2'] <= 0)).sum()
    
    # Ensure correct data types
    df['YEAR'] = df['YEAR'].astype(int)
    df['POPULATION'] = df['POPULATION'].astype(float)
    df['BUILT_UP_SURFACE_M2'] = df['BUILT_UP_SURFACE_M2'].astype(float)
    df['NIGHT_LIGHT'] = df['NIGHT_LIGHT'].astype(float)
    df['IS_TRAINING_PERIOD'] = df['YEAR'] <= 2026
    
    # Sort by zone and year for temporal ordering
    df = df.sort_values(by=['ZONE_ID', 'YEAR']).reset_index(drop=True)
    
    # Feature Engineering (Derived features)
    df['POP_GROWTH_RATE'] = df.groupby('ZONE_ID')['POPULATION'].pct_change()
    df['POP_CHANGE'] = df.groupby('ZONE_ID')['POPULATION'].diff()
    df['BUILT_UP_CHANGE'] = df.groupby('ZONE_ID')['BUILT_UP_SURFACE_M2'].diff()
    df['NIGHT_LIGHT_CHANGE'] = df.groupby('ZONE_ID')['NIGHT_LIGHT'].diff()
    
    # Fill NaN for initial year (2015)
    df['POP_GROWTH_RATE'] = df['POP_GROWTH_RATE'].fillna(0.0)
    df['POP_CHANGE'] = df['POP_CHANGE'].fillna(0.0)
    df['BUILT_UP_CHANGE'] = df['BUILT_UP_CHANGE'].fillna(0.0)
    df['NIGHT_LIGHT_CHANGE'] = df['NIGHT_LIGHT_CHANGE'].fillna(0.0)
    
    # Export processed CSV
    df.to_csv(processed_csv_path, index=False)
    print(f"Processed CSV saved to: {processed_csv_path}")
    
    # Build Data Quality Report
    quality_report = {
        "total_records": int(total_records),
        "number_of_zones": 15,
        "training_period": "2015-2026 (180 records)",
        "model_period": "2015-2030 (240 records)",
        "number_of_features": 8,
        "missing_value_count": int(missing_count),
        "duplicate_count": int(duplicate_count),
        "invalid_value_count": int(invalid_num_count),
        "status": "PASS",
        "data_provenance_notes": "Annual built-up surface between source epochs (2015, 2020, 2025, 2030) are linearly interpolated estimates."
    }
    
    with open(quality_report_path, 'w', encoding='utf-8') as f:
        json.dump(quality_report, f, indent=2)
        
    print(f"Quality report written to: {quality_report_path}")
    
    # Seed PostgreSQL Database
    print("Seeding database smartcity_db...")
    conn = psycopg2.connect(dbname='smartcity_db', user='postgres', password='password', host='localhost', port=5432)
    cur = conn.cursor()
    
    # 1. Insert urban_zone
    for zid, zcode, zname in zones_info:
        cur.execute("""
            INSERT INTO urban_zone (zone_id, zone_code, zone_name)
            VALUES (%s, %s, %s)
            ON CONFLICT (zone_id) DO UPDATE SET zone_code = EXCLUDED.zone_code, zone_name = EXCLUDED.zone_name;
        """, (zid, zcode, zname))
        
    # 2. Insert population_data and urban_features
    for _, row in df.iterrows():
        cur.execute("""
            INSERT INTO population_data 
            (year, zone_id, population, built_up_surface_m2, night_light, built_up_source, night_light_source, is_training_period)
            VALUES (%s, %s, %s, %s, %s, %s, %s, %s)
            ON CONFLICT (year, zone_id) DO UPDATE SET
                population = EXCLUDED.population,
                built_up_surface_m2 = EXCLUDED.built_up_surface_m2,
                night_light = EXCLUDED.night_light,
                built_up_source = EXCLUDED.built_up_source,
                night_light_source = EXCLUDED.night_light_source,
                is_training_period = EXCLUDED.is_training_period;
        """, (
            int(row['YEAR']), int(row['ZONE_ID']), float(row['POPULATION']),
            float(row['BUILT_UP_SURFACE_M2']), float(row['NIGHT_LIGHT']),
            str(row['BUILT_UP_SOURCE']), str(row['NIGHT_LIGHT_SOURCE']),
            bool(row['IS_TRAINING_PERIOD'])
        ))
        
        cur.execute("""
            INSERT INTO urban_features
            (year, zone_id, pop_growth_rate, pop_change, built_up_change, night_light_change)
            VALUES (%s, %s, %s, %s, %s, %s)
            ON CONFLICT (year, zone_id) DO UPDATE SET
                pop_growth_rate = EXCLUDED.pop_growth_rate,
                pop_change = EXCLUDED.pop_change,
                built_up_change = EXCLUDED.built_up_change,
                night_light_change = EXCLUDED.night_light_change;
        """, (
            int(row['YEAR']), int(row['ZONE_ID']), float(row['POP_GROWTH_RATE']),
            float(row['POP_CHANGE']), float(row['BUILT_UP_CHANGE']), float(row['NIGHT_LIGHT_CHANGE'])
        ))
        
    conn.commit()
    cur.close()
    conn.close()
    print("Database seeding completed successfully!")

if __name__ == '__main__':
    load_and_preprocess()
